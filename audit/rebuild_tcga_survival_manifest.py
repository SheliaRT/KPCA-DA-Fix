"""Map local TCGA expression sample IDs to the official TCGA-CDR DSS outcomes.

Only primary-tumor specimens enter the survival cohort audit. Labels follow the
paper's stated extreme thresholds: DSS event within five years is aggressive;
DSS follow-up of at least ten years is indolent; other or missing outcomes are
excluded. The script writes a manifest, not a training dataset, because these
local cohort counts do not yet reproduce the manuscript's stated n=160.

Install the small audit-only dependency with:
    .venv/bin/python -m pip install -r requirements-audit.txt
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CDR_URL = "https://api.gdc.cancer.gov/data/1b5f413e-a8d1-4d10-92eb-7c4ae739ed81"
DEFAULT_CDR = ROOT / "Data/Survival_TCGA_VS_Desmd/TCGA-CDR-SupplementalTableS1.xlsx"
DEFAULT_OUTPUT = ROOT / "audit/tcga_survival_cohort_manifest.json"
EXPRESSION_FILES = {
    "full_1211_sample_matrix": ROOT / "Data/Survival_TCGA_VS_Desmd/TCGA_BRCA_NvT.csv",
    "legacy_154_column_matrix": ROOT / "Data/Survival_TCGA_VS_Desmd/modified_TCGA_BRCA_log2.csv",
    "legacy_250_sample_tvsn_matrix": ROOT / "Data/Survival_TCGA_VS_Desmd/TCGA_BRCA_NvT_v2.csv",
}
DAY = 365.25


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def download_if_missing(path: Path) -> None:
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    request = urllib.request.Request(CDR_URL, headers={"User-Agent": "KPCA-DA-audit/1.0"})
    with urllib.request.urlopen(request, timeout=120) as response, path.open("wb") as output:
        while True:
            block = response.read(1024 * 1024)
            if not block:
                break
            output.write(block)


def sample_id_parts(sample_id: str) -> tuple[str | None, str | None]:
    normalized = str(sample_id).replace(".", "-")
    parts = normalized.split("-")
    if len(parts) < 4 or parts[0] != "TCGA":
        return None, None
    return "-".join(parts[:3]), parts[3][:2]


def load_cdr(path: Path) -> tuple[dict[str, dict[str, Any]], str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "TCGA-CDR" not in workbook.sheetnames:
        raise ValueError("Official TCGA-CDR workbook is missing the TCGA-CDR sheet")
    iterator = workbook["TCGA-CDR"].iter_rows(values_only=True)
    header = next(iterator)
    columns = {str(name): index for index, name in enumerate(header) if name is not None}
    required = {"bcr_patient_barcode", "type", "DSS", "DSS.time"}
    if not required.issubset(columns):
        raise ValueError(f"TCGA-CDR sheet lacks columns: {sorted(required-set(columns))}")
    records = {}
    for row in iterator:
        if row[columns["type"]] != "BRCA":
            continue
        patient = row[columns["bcr_patient_barcode"]]
        if not patient:
            continue
        records[str(patient)] = {
            "dss_event": row[columns["DSS"]],
            "dss_time_days": row[columns["DSS.time"]],
            "vital_status": row[columns["vital_status"]]
            if "vital_status" in columns else None,
            "tumor_status": row[columns["tumor_status"]]
            if "tumor_status" in columns else None,
        }
    workbook.close()
    return records, sha256(path)


def classify_dss(record: dict[str, Any]) -> str:
    event = record.get("dss_event")
    time_days = record.get("dss_time_days")
    if not isinstance(event, (int, float)) or not isinstance(time_days, (int, float)):
        return "missing"
    if event == 1 and time_days <= 5 * DAY:
        return "aggressive"
    if time_days >= 10 * DAY:
        return "indolent"
    return "intermediate_or_excluded"


def sample_headers(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        header = next(csv.reader(handle))
    return [str(sample).strip() for sample in header[1:]]


def audit(cdr_path: Path, output: Path) -> dict[str, Any]:
    download_if_missing(cdr_path)
    cdr, cdr_hash = load_cdr(cdr_path)
    cohorts: dict[str, Any] = {}
    for name, path in EXPRESSION_FILES.items():
        sample_ids = sample_headers(path)
        grouped: dict[str, list[tuple[str, str | None]]] = defaultdict(list)
        type_counts: Counter[str] = Counter()
        for sample_id in sample_ids:
            patient, tissue_code = sample_id_parts(sample_id)
            type_counts[tissue_code or "unparsed"] += 1
            if patient:
                grouped[patient].append((sample_id, tissue_code))
        primary = {patient: entries for patient, entries in grouped.items()
                   if any(tissue == "01" for _, tissue in entries)}
        outcomes: Counter[str] = Counter()
        mapped: list[dict[str, Any]] = []
        missing_cdr: list[str] = []
        duplicate_primary = []
        for patient, entries in sorted(primary.items()):
            primary_ids = [sample_id for sample_id, tissue in entries if tissue == "01"]
            if len(primary_ids) > 1:
                duplicate_primary.append({"patient_id": patient, "sample_ids": primary_ids})
            record = cdr.get(patient)
            if record is None:
                missing_cdr.append(patient)
                category = "missing_cdr_record"
                record = {"dss_event": None, "dss_time_days": None,
                          "vital_status": None, "tumor_status": None}
            else:
                category = classify_dss(record)
            outcomes[category] += 1
            mapped.append({
                "patient_id": patient,
                "primary_tumor_sample_ids": primary_ids,
                **record,
                "paper_threshold_class": category,
            })
        cohorts[name] = {
            "expression_path": str(path.relative_to(ROOT)),
            "expression_sha256": sha256(path),
            "sample_count": len(sample_ids),
            "sample_type_counts": dict(type_counts),
            "unique_patient_ids": len(grouped),
            "primary_tumor_patient_count": len(primary),
            "duplicate_primary_tumor_patient_ids": duplicate_primary,
            "cdr_matched_primary_tumor_patients": len(primary) - len(missing_cdr),
            "cdr_missing_patient_ids": missing_cdr,
            "dss_threshold_counts": dict(outcomes),
            "patient_records": mapped,
        }
    manifest = {
        "generated_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "clinical_source": "TCGA PanCanAtlas TCGA-CDR-SupplementalTableS1.xlsx",
        "clinical_source_url": CDR_URL,
        "clinical_source_sha256": cdr_hash,
        "endpoint_definition": {
            "aggressive": "DSS event=1 and DSS.time <= 5*365.25 days",
            "indolent": "DSS.time >= 10*365.25 days, regardless of later event",
            "otherwise": "intermediate_or_excluded; nonnumeric endpoint values are missing",
            "source_caveat": "The TCGA-CDR notes say DSS is approximated for BRCA, not perfectly observed.",
        },
        "paper_cohort_count": 160,
        "local_expression_cohorts": cohorts,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cdr", type=Path, default=DEFAULT_CDR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = audit(args.cdr, args.output)
    print(json.dumps({
        "output": str(args.output),
        "cohorts": {
            name: {key: value[key] for key in (
                "sample_count", "sample_type_counts", "primary_tumor_patient_count",
                "cdr_matched_primary_tumor_patients", "dss_threshold_counts",
            )}
            for name, value in result["local_expression_cohorts"].items()
        },
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
